"""Build one Excel workbook of company, dividend and earnings data.

Works offline from what build_universe.py and fetch_fundamentals.py saved:
    <out>/_state/rows.json               symbol matching and listing status
    <out>/_state/quotesummary/*.json     Yahoo profile / dividend / earnings data
    <out>/prices/<SYMBOL>/1d.csv         past dividends (Dividends column)

Sheets: Read Me, Overview, Dividend History, Dividends by Year,
Earnings Quarterly, Earnings Annual, Analyst Estimates.

Usage:
    python scripts/build_fundamentals_workbook.py Tickers.xlsx --out Tickers_out
"""

import argparse
import datetime as dt
import json
import os
import re
import sys

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TODAY = dt.date.today()
FONT = Font(name="Arial", size=10)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="133F45")
GROUP_FONT = Font(name="Arial", size=10, bold=True, color="133F45")

MONEY = "#,##0.00"
MONEY4 = "#,##0.0000"
INT = "#,##0"
MILLIONS = "#,##0.0"
PCT = "0.00%"
DATE = "yyyy-mm-dd"
RATIO = "0.00"


def raw(d, *path):
    """Walk nested Yahoo JSON; unwrap {'raw': x} and treat {} as missing."""
    for p in path:
        if isinstance(d, list):
            d = d[p] if isinstance(p, int) and len(d) > p else None
        elif isinstance(d, dict):
            d = d.get(p)
        else:
            return None
        if d is None:
            return None
    if isinstance(d, dict):
        return d.get("raw") if "raw" in d else None
    return d


def to_date(ts):
    if ts in (None, "", 0):
        return None
    if isinstance(ts, str):
        try:
            return dt.date.fromisoformat(ts[:10])
        except ValueError:
            return None
    return dt.datetime.fromtimestamp(ts, dt.timezone.utc).date()


def millions(x):
    return x / 1e6 if isinstance(x, (int, float)) else None


def pct_from_percent(x):
    # A few Yahoo fields (e.g. fiveYearAvgDividendYield) are already in percent.
    return x / 100 if isinstance(x, (int, float)) else None


def load_inputs(xlsx, out):
    src = pd.read_excel(xlsx)
    state = json.load(open(os.path.join(out, "_state", "rows.json")))
    qs_dir = os.path.join(out, "_state", "quotesummary")
    qs = {}
    for f in os.listdir(qs_dir):
        if f.endswith(".json"):
            qs[f[:-5]] = json.load(open(os.path.join(qs_dir, f))) or {}
    divs = {}
    for sym in os.listdir(os.path.join(out, "prices")):
        p = os.path.join(out, "prices", sym, "1d.csv")
        if os.path.exists(p):
            d = pd.read_csv(p, usecols=["Date", "Dividends"], parse_dates=["Date"])
            d = d[d["Dividends"] > 0]
            divs[sym] = [(r.Date.date(), float(r.Dividends)) for r in d.itertuples()]
    return src, state, qs, divs


def price_to_book(st, sd, ks, fd):
    """Yahoo divides a pence price by a pounds book value for London stocks
    (giving ~100x too high), so compute it: price in whole units / book value.
    Left blank when the company reports in a different currency than it trades."""
    price = raw(fd, "currentPrice") or raw(sd, "previousClose")
    book = raw(ks, "bookValue")
    if not price or not book or book <= 0:
        return None
    minor = (st.get("Yahoo Currency Code") or "") in ("GBp", "GBX", "ZAc", "ILA")
    if fd.get("financialCurrency") and fd["financialCurrency"] != st.get("Currency"):
        return None
    return price / (100 if minor else 1) / book


def overview_rows(src, state, qs, divs):
    rows = []
    year_ago = TODAY - dt.timedelta(days=365)
    for _, r in src.iterrows():
        st = state.get(r["ISIN"], {})
        sym = st.get("Yahoo Symbol") or None
        q = qs.get(sym, {}) if sym else {}
        hist = divs.get(sym, []) if sym else []
        ap, sd, ks, fd, ce = (q.get(k, {}) or {} for k in
                              ("assetProfile", "summaryDetail", "defaultKeyStatistics", "financialData", "calendarEvents"))
        fp = q.get("fundProfile") or {}
        eh = (q.get("earningsHistory") or {}).get("history") or []
        last_q = eh[-1] if eh else {}
        ex_date = to_date(raw(ce, "exDividendDate") or raw(sd, "exDividendDate"))
        pay_date = to_date(raw(ce, "dividendDate"))
        earn_dates = (ce.get("earnings") or {}).get("earningsDate") or []
        next_earn = to_date(raw(earn_dates, 0))
        is_fund = bool(fp) or st.get("Instrument Type") in ("ETF", "Fund / ETP")
        address = ", ".join(x for x in (ap.get("city"), ap.get("state"), ap.get("country")) if x)

        rows.append({
            # Identity
            "No.": r["No."], "ISIN": r["ISIN"], "Ticker": r["Ticker"], "Instrument Name": r["Instrument Name"],
            "Yahoo Symbol": sym, "Yahoo Name": st.get("Yahoo Name"), "Status": st.get("Status"),
            "Instrument Type": st.get("Instrument Type"), "Exchange": st.get("Exchange"),
            "Currency": st.get("Currency"), "Price Unit": st.get("Price Unit"),
            # Company
            "Sector": ap.get("sectorDisp") or ap.get("sector"), "Industry": ap.get("industryDisp") or ap.get("industry"),
            "HQ City": ap.get("city"), "HQ Country": ap.get("country"), "HQ Address": address or None,
            "Website": ap.get("website"), "Employees": ap.get("fullTimeEmployees"),
            # Valuation
            "Market Cap (millions, Currency)": millions(raw(sd, "marketCap")),
            "Trailing P/E": raw(sd, "trailingPE"), "Forward P/E": raw(sd, "forwardPE"),
            "Price/Book": price_to_book(st, sd, ks, fd), "Beta": raw(sd, "beta"),
            "52w Low (Price Unit)": raw(sd, "fiftyTwoWeekLow"), "52w High (Price Unit)": raw(sd, "fiftyTwoWeekHigh"),
            # Dividends - Yahoo's current view
            "Annual Dividend Rate (Currency)": raw(sd, "dividendRate"),
            "Dividend Yield": raw(sd, "dividendYield") or raw(sd, "yield"),
            "Payout Ratio": raw(sd, "payoutRatio"),
            "5y Avg Dividend Yield": pct_from_percent(raw(sd, "fiveYearAvgDividendYield")),
            "Latest Announced Ex-Div Date": ex_date,
            "Latest Announced Payment Date": pay_date,
            "Ex-Div Date Upcoming?": ("Yes" if ex_date and ex_date >= TODAY else "No") if ex_date else None,
            # Dividends - from downloaded history
            "Dividends on Record": len(hist) or None,
            "First Ex-Date on Record": hist[0][0] if hist else None,
            "Last Ex-Date on Record": hist[-1][0] if hist else None,
            "Last Dividend (Price Unit)": hist[-1][1] if hist else None,
            "Last 12m Dividends (Price Unit)": (sum(a for d, a in hist if d > year_ago) or None) if hist else None,
            # Earnings
            "Next Earnings Date": next_earn,
            "Next Earnings Date Is Estimate?": ("Yes" if (ce.get("earnings") or {}).get("isEarningsDateEstimate") else "No") if next_earn else None,
            "Next Qtr EPS Estimate": raw(ce, "earnings", "earningsAverage"),
            "Next Qtr Revenue Estimate (millions)": millions(raw(ce, "earnings", "revenueAverage")),
            "Last Reported Qtr End": to_date(raw(last_q, "quarter")),
            "Last Qtr EPS Actual": raw(last_q, "epsActual"),
            "Last Qtr EPS Estimate": raw(last_q, "epsEstimate"),
            "Last Qtr EPS Surprise": raw(last_q, "surprisePercent"),
            "Trailing EPS": raw(ks, "trailingEps"), "Forward EPS": raw(ks, "forwardEps"),
            "Revenue TTM (millions)": millions(raw(fd, "totalRevenue")),
            "Profit Margin": raw(fd, "profitMargins"), "Revenue Growth (YoY)": raw(fd, "revenueGrowth"),
            "Earnings Growth (YoY)": raw(fd, "earningsGrowth"), "Return on Equity": raw(fd, "returnOnEquity"),
            "Debt/Equity (%)": raw(fd, "debtToEquity"),
            "Financials Currency": fd.get("financialCurrency"),
            "Analyst Rating": (fd.get("recommendationKey") or "none").replace("_", " ").title().replace("None", "") or None,
            "Analyst Count": raw(fd, "numberOfAnalystOpinions"),
            "Target Price Mean (Price Unit)": raw(fd, "targetMeanPrice"),
            # Funds
            "Fund Family": (fp.get("family") or ks.get("fundFamily")) if is_fund else None,
            "Fund Category": (fp.get("categoryName") or ks.get("category")) if is_fund else None,
            "Fund Legal Type": (fp.get("legalType") or ks.get("legalType")) if is_fund else None,
            "Fund Total Assets (millions)": millions(raw(ks, "totalAssets") or raw(sd, "totalAssets")) if is_fund else None,
            "Expense Ratio": (raw(fp, "feesExpensesInvestment", "annualReportExpenseRatio")
                              or raw(fp, "feesExpensesInvestment", "netExpRatio")
                              or raw(ks, "annualReportExpenseRatio")) if is_fund else None,
            "Fund Inception": to_date(raw(ks, "fundInceptionDate")) if is_fund else None,
            "Profile Data": ("Yes" if q else "No") if sym else None,
        })
    return rows


OVERVIEW_GROUPS = [
    ("Instrument", ["No.", "ISIN", "Ticker", "Instrument Name", "Yahoo Symbol", "Yahoo Name", "Status",
                    "Instrument Type", "Exchange", "Currency", "Price Unit"]),
    ("Company", ["Sector", "Industry", "HQ City", "HQ Country", "HQ Address", "Website", "Employees"]),
    ("Valuation", ["Market Cap (millions, Currency)", "Trailing P/E", "Forward P/E", "Price/Book", "Beta",
                   "52w Low (Price Unit)", "52w High (Price Unit)"]),
    ("Dividends - Yahoo's current figures", ["Annual Dividend Rate (Currency)", "Dividend Yield", "Payout Ratio",
                                             "5y Avg Dividend Yield", "Latest Announced Ex-Div Date",
                                             "Latest Announced Payment Date", "Ex-Div Date Upcoming?"]),
    ("Dividends - from price history", ["Dividends on Record", "First Ex-Date on Record", "Last Ex-Date on Record",
                                        "Last Dividend (Price Unit)", "Last 12m Dividends (Price Unit)"]),
    ("Earnings", ["Next Earnings Date", "Next Earnings Date Is Estimate?", "Next Qtr EPS Estimate",
                  "Next Qtr Revenue Estimate (millions)", "Last Reported Qtr End", "Last Qtr EPS Actual",
                  "Last Qtr EPS Estimate", "Last Qtr EPS Surprise", "Trailing EPS", "Forward EPS",
                  "Revenue TTM (millions)", "Profit Margin", "Revenue Growth (YoY)", "Earnings Growth (YoY)",
                  "Return on Equity", "Debt/Equity (%)", "Financials Currency", "Analyst Rating", "Analyst Count",
                  "Target Price Mean (Price Unit)"]),
    ("Funds / ETFs", ["Fund Family", "Fund Category", "Fund Legal Type", "Fund Total Assets (millions)",
                      "Expense Ratio", "Fund Inception", "Profile Data"]),
]

FORMATS = {
    "Employees": INT, "Market Cap (millions, Currency)": MILLIONS, "Trailing P/E": RATIO, "Forward P/E": RATIO,
    "Price/Book": RATIO, "Beta": RATIO, "52w Low (Price Unit)": MONEY, "52w High (Price Unit)": MONEY,
    "Annual Dividend Rate (Currency)": MONEY4, "Dividend Yield": PCT, "Payout Ratio": PCT,
    "5y Avg Dividend Yield": PCT, "Latest Announced Ex-Div Date": DATE, "Latest Announced Payment Date": DATE,
    "Dividends on Record": INT, "First Ex-Date on Record": DATE, "Last Ex-Date on Record": DATE,
    "Last Dividend (Price Unit)": MONEY4, "Last 12m Dividends (Price Unit)": MONEY4,
    "Next Earnings Date": DATE, "Next Qtr EPS Estimate": MONEY4, "Next Qtr Revenue Estimate (millions)": MILLIONS,
    "Last Reported Qtr End": DATE, "Last Qtr EPS Actual": MONEY4, "Last Qtr EPS Estimate": MONEY4,
    "Last Qtr EPS Surprise": PCT, "Trailing EPS": MONEY4, "Forward EPS": MONEY4, "Revenue TTM (millions)": MILLIONS,
    "Profit Margin": PCT, "Revenue Growth (YoY)": PCT, "Earnings Growth (YoY)": PCT, "Return on Equity": PCT,
    "Debt/Equity (%)": RATIO, "Analyst Count": INT, "Target Price Mean (Price Unit)": MONEY,
    "Fund Total Assets (millions)": MILLIONS, "Expense Ratio": PCT, "Fund Inception": DATE,
}


def style_header(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_table(ws, name, first_row, ncols, nrows):
    ref = f"A{first_row}:{get_column_letter(ncols)}{first_row + max(nrows, 1)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)


def write_sheet(wb, title, columns, rows, formats=None, widths=None, table_name=None, groups=None, note=None):
    """Write a list of dicts as an Excel table. Values may be formulas ('=...')
    containing {r} for the row number. Formulas refer to columns by letter, so
    the column order passed in must match them."""
    ws = wb.create_sheet(title)
    formats, widths = formats or {}, widths or {}
    header_row = 1
    if note:
        ws.cell(1, 1, note).font = Font(name="Arial", size=10, italic=True)
        header_row = 2
    if groups:
        col = 1
        for label, cols in groups:
            ws.cell(header_row, col, label).font = GROUP_FONT
            col += len(cols)
        header_row += 1
    for j, c in enumerate(columns, 1):
        style_header(ws.cell(header_row, j, c))
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, max(10, min(len(c) + 2, 22)))
    for i, row in enumerate(rows, header_row + 1):
        for j, c in enumerate(columns, 1):
            v = row.get(c)
            if isinstance(v, str) and v.startswith("="):
                v = v.format(r=i)
            if isinstance(v, float) and v != v:  # NaN
                v = None
            cell = ws.cell(i, j, v)
            cell.font = FONT
            if c in formats:
                cell.number_format = formats[c]
    ws.row_dimensions[header_row].height = 42
    add_table(ws, table_name or re.sub(r"\W", "", title), header_row, len(columns), len(rows))
    ws.freeze_panes = ws.cell(header_row + 1, 1 if title != "Overview" else 6)
    return ws


def dividend_history_rows(overview, divs):
    seen, rows = set(), []
    for o in overview:
        sym = o["Yahoo Symbol"]
        if not sym or sym in seen:
            continue
        seen.add(sym)
        for d, a in divs.get(sym, []):
            rows.append({"Yahoo Symbol": sym, "Name": o["Yahoo Name"], "Ex-Dividend Date": d,
                         "Amount (Price Unit)": a, "Price Unit": o["Price Unit"], "Year": d.year})
    return rows


def earnings_quarterly_rows(overview, qs):
    rows, seen = [], set()
    for o in overview:
        sym = o["Yahoo Symbol"]
        if not sym or sym in seen:
            continue
        seen.add(sym)
        q = qs.get(sym) or {}
        e = q.get("earnings") or {}
        fin = {x.get("date"): x for x in (e.get("financialsChart") or {}).get("quarterly") or []}
        ccy = (q.get("financialData") or {}).get("financialCurrency") or e.get("financialCurrency")
        for x in (e.get("earningsChart") or {}).get("quarterly") or []:
            f = fin.get(x.get("date"), {})
            rows.append({
                "Yahoo Symbol": sym, "Name": o["Yahoo Name"], "Quarter": x.get("calendarQuarter") or x.get("date"),
                "Fiscal Quarter": x.get("fiscalQuarter"), "Period End": to_date(raw(x, "periodEndDate")),
                "Reported": to_date(raw(x, "reportedDate")), "EPS Actual": raw(x, "actual"),
                "EPS Estimate": raw(x, "estimate"),
                "EPS Difference": '=IF(OR(F{r}="",G{r}=""),"",F{r}-G{r})',
                "EPS Surprise %": '=IF(OR(F{r}="",G{r}="",G{r}=0),"",(F{r}-G{r})/ABS(G{r}))',
                "Revenue (millions)": millions(raw(f, "revenue")), "Net Income (millions)": millions(raw(f, "earnings")),
                "Currency": ccy,
            })
    return rows


def earnings_annual_rows(overview, qs):
    rows, seen = [], set()
    for o in overview:
        sym = o["Yahoo Symbol"]
        if not sym or sym in seen:
            continue
        seen.add(sym)
        q = qs.get(sym) or {}
        ccy = (q.get("financialData") or {}).get("financialCurrency")
        for x in ((q.get("earnings") or {}).get("financialsChart") or {}).get("yearly") or []:
            rows.append({"Yahoo Symbol": sym, "Name": o["Yahoo Name"], "Fiscal Year": x.get("date"),
                         "Revenue (millions)": millions(raw(x, "revenue")),
                         "Net Income (millions)": millions(raw(x, "earnings")),
                         "Profit Margin": '=IF(OR(D{r}="",E{r}="",D{r}=0),"",E{r}/D{r})', "Currency": ccy})
    return rows


PERIODS = {"0q": "Current Quarter", "+1q": "Next Quarter", "0y": "Current Year", "+1y": "Next Year"}


def estimate_rows(overview, qs):
    rows, seen = [], set()
    for o in overview:
        sym = o["Yahoo Symbol"]
        if not sym or sym in seen:
            continue
        seen.add(sym)
        for t in ((qs.get(sym) or {}).get("earningsTrend") or {}).get("trend") or []:
            if t.get("period") not in PERIODS:
                continue
            ee, re_ = t.get("earningsEstimate") or {}, t.get("revenueEstimate") or {}
            if raw(ee, "avg") is None and raw(re_, "avg") is None:
                continue
            rows.append({
                "Yahoo Symbol": sym, "Name": o["Yahoo Name"], "Period": PERIODS[t["period"]],
                "Period End": to_date(t.get("endDate")), "EPS Estimate (avg)": raw(ee, "avg"),
                "EPS Low": raw(ee, "low"), "EPS High": raw(ee, "high"), "EPS Year Ago": raw(ee, "yearAgoEps"),
                "EPS Growth": raw(ee, "growth"), "EPS Analysts": raw(ee, "numberOfAnalysts"),
                "Revenue Estimate (millions)": millions(raw(re_, "avg")),
                "Revenue Year Ago (millions)": millions(raw(re_, "yearAgoRevenue")),
                "Revenue Growth": raw(re_, "growth"), "Revenue Analysts": raw(re_, "numberOfAnalysts"),
                "Currency": ee.get("earningsCurrency") or re_.get("revenueCurrency"),
            })
    return rows


README = [
    ("Tickers_4Oct26 - Company, dividend and earnings data", None),
    (f"Source: Yahoo Finance, downloaded {TODAY:%d %b %Y}. Yahoo data can contain gaps and errors; "
     "check anything important against company filings.", None),
    ("", None),
    ("Sheets", "bold"),
    ("Overview", "One row per row of your list: sector, HQ, valuation, dividends, earnings, fund details."),
    ("Dividend History", "Every past dividend on record: ex-dividend date and amount (one row per payment)."),
    ("Dividends by Year", "Total dividends per calendar year per instrument (formulas over Dividend History)."),
    ("Earnings Quarterly", "Last 4 reported quarters: EPS actual vs analyst estimate, revenue, net income."),
    ("Earnings Annual", "Last 4 fiscal years: revenue, net income, profit margin."),
    ("Analyst Estimates", "Consensus EPS and revenue for the current/next quarter and current/next year."),
    ("", None),
    ("Units - please read", "bold"),
    ("Price Unit", "Dividend history, 52-week range and prices are in the instrument's Price Unit. "
                   "For London that is often PENCE (GBp): 17.0 means 17p, not GBP 17."),
    ("Currency", "Yahoo's Annual Dividend Rate and Market Cap are in whole currency units, "
                 "e.g. GBP 0.17 for the same 17p dividend. Target Price is in the Price Unit."),
    ("Price/Book", "Calculated here (price / book value per share) because Yahoo's figure is ~100x too high "
                   "for pence-quoted stocks. Blank when the company reports in another currency."),
    ("Financials Currency", "Earnings, revenue and EPS are in the currency the company reports in, "
                            "which can differ from the trading currency (e.g. a London stock reporting in USD)."),
    ("(millions)", "Columns marked (millions) are divided by 1,000,000."),
    ("", None),
    ("Limits of the data", "bold"),
    ("Payment dates", "Yahoo only gives the payment date for the latest announced dividend, not for past ones."),
    ("Future dividends", "Shown only once a company has announced them ('Ex-Div Date Upcoming?' = Yes)."),
    ("Ex-div dates", "Dividend History uses the exchange's ex-date. Yahoo occasionally misses or duplicates one."),
    ("Earnings history", "Yahoo's free data covers only the last 4 quarters and 4 years."),
    ("Funds / ETFs", "No sector, earnings or HQ; fund family, category, size and fees where Yahoo has them."),
    ("Not found", "Instruments not on Yahoo (mostly delisted) appear in Overview with no data."),
]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--out", required=True)
    ap.add_argument("--name", default="Tickers_4Oct26_Fundamentals.xlsx")
    ap.add_argument("--first-year", type=int, default=2000)
    args = ap.parse_args()

    src, state, qs, divs = load_inputs(args.xlsx, args.out)
    overview = overview_rows(src, state, qs, divs)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    for i, (a, b) in enumerate(README, 1):
        c = ws.cell(i, 1, a)
        c.font = Font(name="Arial", size=12 if i == 1 else 10, bold=(i == 1 or b == "bold"))
        if b and b != "bold":
            ws.cell(i, 2, b).font = FONT
            ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name="Arial", size=10, bold=True)

    cols = [c for _, cs in OVERVIEW_GROUPS for c in cs]
    write_sheet(wb, "Overview", cols, overview, FORMATS, groups=OVERVIEW_GROUPS,
                widths={"Instrument Name": 26, "Yahoo Name": 30, "Status": 22, "Industry": 26, "HQ Address": 30,
                        "Website": 28, "Sector": 20, "Fund Category": 24, "Fund Family": 28})

    hist = dividend_history_rows(overview, divs)
    write_sheet(wb, "Dividend History", ["Yahoo Symbol", "Name", "Ex-Dividend Date", "Amount (Price Unit)",
                                         "Price Unit", "Year"], hist,
                {"Ex-Dividend Date": DATE, "Amount (Price Unit)": MONEY4}, {"Name": 34, "Price Unit": 18},
                table_name="DividendHistory")

    # Dividends by Year: SUMIFS over the Dividend History table.
    n = len(hist) + 1
    years = list(range(args.first_year, TODAY.year + 1))
    by_year, seen = [], set()
    for o in overview:
        sym = o["Yahoo Symbol"]
        if sym and sym not in seen and divs.get(sym):
            seen.add(sym)
            row = {"Yahoo Symbol": sym, "Name": o["Yahoo Name"], "Price Unit": o["Price Unit"]}
            for y in years:
                row[str(y)] = (f"=SUMIFS('Dividend History'!$D$2:$D${n},'Dividend History'!$A$2:$A${n},$A{{r}},"
                               f"'Dividend History'!$F$2:$F${n},{y})")
            by_year.append(row)
    write_sheet(wb, "Dividends by Year", ["Yahoo Symbol", "Name", "Price Unit"] + [str(y) for y in years], by_year,
                {str(y): "#,##0.00;-#,##0.00;-" for y in years}, {"Name": 30, "Price Unit": 16},
                table_name="DividendsByYear",
                note=f"Total dividends per calendar year, in each instrument's Price Unit. "
                     f"Years before {args.first_year} are in Dividend History.")

    write_sheet(wb, "Earnings Quarterly", ["Yahoo Symbol", "Name", "Quarter", "Fiscal Quarter", "Period End",
                                           "EPS Actual", "EPS Estimate", "EPS Difference", "EPS Surprise %",
                                           "Reported", "Revenue (millions)", "Net Income (millions)", "Currency"],
                earnings_quarterly_rows(overview, qs),
                {"Period End": DATE, "Reported": DATE, "EPS Actual": MONEY4, "EPS Estimate": MONEY4,
                 "EPS Difference": MONEY4, "EPS Surprise %": PCT, "Revenue (millions)": MILLIONS,
                 "Net Income (millions)": MILLIONS}, {"Name": 30}, table_name="EarningsQuarterly")

    write_sheet(wb, "Earnings Annual", ["Yahoo Symbol", "Name", "Fiscal Year", "Revenue (millions)",
                                        "Net Income (millions)", "Profit Margin", "Currency"],
                earnings_annual_rows(overview, qs),
                {"Revenue (millions)": MILLIONS, "Net Income (millions)": MILLIONS, "Profit Margin": PCT},
                {"Name": 30}, table_name="EarningsAnnual")

    est_cols = ["Yahoo Symbol", "Name", "Period", "Period End", "EPS Estimate (avg)", "EPS Low", "EPS High",
                "EPS Year Ago", "EPS Growth", "EPS Analysts", "Revenue Estimate (millions)",
                "Revenue Year Ago (millions)", "Revenue Growth", "Revenue Analysts", "Currency"]
    write_sheet(wb, "Analyst Estimates", est_cols, estimate_rows(overview, qs),
                {"Period End": DATE, "EPS Estimate (avg)": MONEY4, "EPS Low": MONEY4, "EPS High": MONEY4,
                 "EPS Year Ago": MONEY4, "EPS Growth": PCT, "Revenue Estimate (millions)": MILLIONS,
                 "Revenue Year Ago (millions)": MILLIONS, "Revenue Growth": PCT}, {"Name": 30},
                table_name="AnalystEstimates")

    wb.calculation.fullCalcOnLoad = True
    dst = os.path.join(args.out, args.name)
    wb.save(dst)
    print(f"Wrote {dst}")
    print(f"Overview {len(overview)}, dividends {len(hist)}, by-year {len(by_year)}")


if __name__ == "__main__":
    sys.exit(main())
