"""Resolve a ticker list to Yahoo symbols, check listing status, and save daily history.

Reads a spreadsheet with ISIN / Ticker / Instrument Name columns, and for each row:

1. Matches it to a Yahoo symbol. Searches Yahoo by ISIN and prefers the plain
   ticker (US listing) or ticker + ".L" (London), ordered by the ISIN's country.
   Foreign secondary listings (e.g. Mexico) are only used as a last resort.
2. Downloads the full daily history in one request per symbol. The same
   response carries listing status, currency (GBP vs GBp pence), instrument
   type and exchange, so the check costs nothing extra.
3. Writes those details as new columns in a copy of the spreadsheet.

Calls Yahoo's search and chart endpoints directly (the same ones yfinance
uses) because they need no cookie/crumb, which is the step Yahoo throttles
hardest, and so HTTP 429 rate limits are visible and retried rather than
turning into silently empty data.

Progress is checkpointed to <out>/_state/rows.json; re-running resumes.

Usage:
    python scripts/build_universe.py Tickers.xlsx --out Tickers_out
"""

import argparse
import datetime as dt
import json
import logging
import os
import re
import sys
import time
from copy import copy

import openpyxl
import pandas as pd
import requests
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn

logger = logging.getLogger("build_universe")

SEARCH_URL = "https://query2.finance.yahoo.com/v1/finance/search"
CHART_URL = "https://query2.finance.yahoo.com/v8/finance/chart/{}"
NINETY_NINE_YEARS = 3122064000  # same lookback yfinance uses for period="max"

# ISIN countries whose instruments this list most likely holds in London.
LONDON_FIRST = {"GB", "JE", "GG", "IM", "IE", "LU", "XS"}

STOPWORDS = {"plc", "inc", "corp", "corporation", "ltd", "limited", "the", "group", "holdings",
             "holding", "co", "company", "sa", "nv", "ag", "se", "etf", "ucits", "class", "ord",
             "and", "of", "shares", "share", "adr", "ads", "trust", "fund", "acc", "dist", "usd",
             "gbp", "eur"}

INSTRUMENT_TYPES = {"EQUITY": "Stock", "ETF": "ETF", "MUTUALFUND": "Fund / ETP", "INDEX": "Index",
                    "CURRENCY": "Currency", "CRYPTOCURRENCY": "Crypto", "FUTURE": "Future"}

# Yahoo quotes some markets in minor units.
MINOR_UNITS = {"GBp": ("GBP", "Pence (GBp)"), "GBX": ("GBP", "Pence (GBp)"),
               "ZAc": ("ZAR", "Cents (ZAc)"), "ILA": ("ILS", "Agorot (ILA)")}


MAJOR_UNITS = {"GBP": "Pounds (GBP)", "USD": "US dollars", "EUR": "Euros", "CAD": "Canadian dollars",
               "HKD": "Hong Kong dollars", "CHF": "Swiss francs", "AUD": "Australian dollars",
               "JPY": "Japanese yen", "SEK": "Swedish kronor", "NOK": "Norwegian kroner",
               "DKK": "Danish kroner", "ZAR": "Rand"}


class Yahoo:
    """Throttled client for Yahoo's public endpoints with 429 backoff."""

    def __init__(self, min_interval=2.0):
        self.s = requests.Session()
        self.s.headers["User-Agent"] = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                                        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
        self.min_interval = min_interval
        self.last = 0.0
        self.requests = 0

    def get(self, url, params):
        backoff = 60
        for attempt in range(8):
            wait = self.last + self.min_interval - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            self.last = time.monotonic()
            self.requests += 1
            try:
                r = self.s.get(url, params=params, timeout=30)
            except requests.RequestException as e:
                logger.warning("network error %s, retrying in 30s", e)
                time.sleep(30)
                continue
            if r.status_code == 429 or r.status_code >= 500:
                logger.warning("HTTP %s, backing off %ds", r.status_code, backoff)
                time.sleep(backoff)
                backoff = min(backoff * 2, 900)
                continue
            return r
        raise RuntimeError(f"gave up on {url} after repeated rate limiting")

    def search(self, query):
        r = self.get(SEARCH_URL, {"q": query, "quotesCount": 10, "newsCount": 0, "listsCount": 0})
        if r.status_code != 200:
            return []
        return [{"symbol": q.get("symbol"), "exchange": q.get("exchange"),
                 "quoteType": q.get("quoteType"), "name": q.get("longname") or q.get("shortname")}
                for q in r.json().get("quotes", []) if q.get("symbol")]

    def daily(self, symbol):
        """Full daily history. Returns (meta, DataFrame) or (None, None) if not found."""
        now = int(time.time())
        r = self.get(CHART_URL.format(symbol), {"period1": now - NINETY_NINE_YEARS, "period2": now,
                                                "interval": "1d", "events": "div,splits",
                                                "includeAdjustedClose": "true"})
        if r.status_code == 404:
            return None, None
        try:
            res = r.json()["chart"]["result"][0]
        except (ValueError, KeyError, IndexError, TypeError):
            return None, None
        meta = res.get("meta", {})
        ts = res.get("timestamp")
        if not ts:
            return meta, pd.DataFrame()
        q = res["indicators"]["quote"][0]
        tz = meta.get("exchangeTimezoneName") or "UTC"
        idx = pd.to_datetime(ts, unit="s", utc=True).tz_convert(tz).date
        df = pd.DataFrame({"Open": q.get("open"), "High": q.get("high"), "Low": q.get("low"),
                           "Close": q.get("close"), "Volume": q.get("volume")}, index=pd.Index(idx, name="Date"))
        adj = res["indicators"].get("adjclose")
        df.insert(4, "Adj Close", adj[0]["adjclose"] if adj else df["Close"])
        df["Dividends"] = 0.0
        df["Stock Splits"] = 0.0
        events = res.get("events", {})
        for e in events.get("dividends", {}).values():
            d = pd.Timestamp(e["date"], unit="s", tz="UTC").tz_convert(tz).date()
            if d in df.index:
                df.loc[d, "Dividends"] = e["amount"]
        for e in events.get("splits", {}).values():
            d = pd.Timestamp(e["date"], unit="s", tz="UTC").tz_convert(tz).date()
            if d in df.index:
                df.loc[d, "Stock Splits"] = e["numerator"] / e["denominator"]
        df = df.dropna(subset=["Open", "High", "Low", "Close"], how="all")
        df = df[~df.index.duplicated(keep="last")]
        return meta, df


def normalize_ticker(t):
    t = str(t).strip().upper().rstrip(".")
    return re.sub(r"[./]", "-", t)


def name_tokens(s):
    return {w for w in re.findall(r"[a-z0-9]+", (s or "").lower()) if w not in STOPWORDS and len(w) > 1}


def names_match(ours, yahoo):
    ta, tb = name_tokens(ours), name_tokens(yahoo)
    if ta & tb:
        return True
    # Renamed after a merger: "ATAI Life Sciences" vs "AtaiBeckley Inc."
    if any(len(a) >= 4 and b.startswith(a) for a in ta for b in tb):
        return True
    # Abbreviations: "ABF" vs "Associated British Foods plc"
    initials = "".join(w[0] for w in re.findall(r"[a-z0-9]+", (yahoo or "").lower()) if w not in STOPWORDS)
    compact = re.sub(r"[^a-z0-9]", "", (ours or "").lower())
    return len(compact) >= 2 and initials.startswith(compact)


def resolve_row(yh, isin, ticker, name):
    """Return dict with chosen symbol, how it was matched, and its daily data."""
    t = normalize_ticker(ticker)
    country = isin[:2]
    prefs = [t + ".L", t] if country in LONDON_FIRST else [t, t + ".L"]
    quotes = yh.search(isin)
    syms = [q["symbol"] for q in quotes]
    notes = []
    tried = {}

    def fetch(sym):
        if sym not in tried:
            tried[sym] = yh.daily(sym)
        return tried[sym]

    # a. ISIN search confirms the plain ticker or the London listing.
    for c in prefs:
        if c in syms:
            meta, df = fetch(c)
            if meta is not None:
                return dict(symbol=c, method="ISIN match", quotes=quotes, notes=notes, meta=meta, df=df)

    if quotes:
        notes.append("Yahoo ISIN search returned: " + ", ".join(syms[:4]))
    else:
        notes.append("ISIN not found on Yahoo (may be an old ISIN after a corporate action)")

    def isin_other_exchange():
        # ISIN search found the same ticker on its home exchange (e.g. AF.PA,
        # ASML.AS). Skip foreign secondary listings like Mexico.
        for q in quotes:
            if q["symbol"].split(".")[0].replace("-", "") == t.replace("-", "") and not q["symbol"].endswith(".MX"):
                meta, df = fetch(q["symbol"])
                if meta is not None:
                    return dict(symbol=q["symbol"], method="ISIN match (other exchange)", quotes=quotes, notes=notes, meta=meta, df=df)

    def ticker_and_name():
        # Ticker exists on the preferred exchange and the name agrees.
        for c in prefs:
            meta, df = fetch(c)
            if meta is not None and df is not None and len(df) and names_match(name, meta.get("longName") or meta.get("shortName")):
                return dict(symbol=c, method="Ticker + name match", quotes=quotes, notes=notes, meta=meta, df=df)

    # b/c. UK/Irish/Jersey-type ISINs are held in London, so the London line
    # wins over e.g. a Swiss listing of the same ETF. Elsewhere the ISIN's own
    # home listing wins.
    steps = (ticker_and_name, isin_other_exchange) if country in LONDON_FIRST else (isin_other_exchange, ticker_and_name)
    for step in steps:
        found = step()
        if found:
            return found

    # d. Ticker exists on the ISIN's home exchange (US or London only) but the
    #    name doesn't obviously match - keep it, flag it. Not done for other
    #    countries: "MARA" in the US is not Marubeni, and "GOOD" is not GOOD.L.
    if country == "US" or country in LONDON_FIRST:
        meta, df = fetch(prefs[0])
        if meta is not None and df is not None and len(df):
            notes.append("Name differs from Yahoo's - please verify")
            return dict(symbol=prefs[0], method="Ticker only - verify", quotes=quotes, notes=notes, meta=meta, df=df)

    # e. Fall back to whatever the ISIN search found.
    for q in quotes:
        meta, df = fetch(q["symbol"])
        if meta is not None and df is not None and len(df):
            notes.append("Ticker not found on Yahoo; using ISIN result - please verify")
            return dict(symbol=q["symbol"], method="ISIN only - verify", quotes=quotes, notes=notes, meta=meta, df=df)

    return dict(symbol=None, method="Not found", quotes=quotes, notes=notes, meta=None, df=None)


def describe(res, today):
    meta, df = res["meta"], res["df"]
    out = {"Yahoo Symbol": res["symbol"] or "", "Match Method": res["method"]}
    if not res["symbol"]:
        out["Status"] = "Not found on Yahoo"
        out["Notes"] = "; ".join(res["notes"])
        return out
    last = df.index[-1] if len(df) else None
    days = (today - last).days if last else None
    if last is None:
        status = "No price data"
    elif days <= 10:
        status = "Listed"
    else:
        status = f"Not trading since {last:%Y-%m-%d} (likely delisted/suspended)"
    yccy = meta.get("currency") or ""
    ccy, unit = MINOR_UNITS.get(yccy, (yccy.upper(), MAJOR_UNITS.get(yccy, yccy)))
    itype = meta.get("instrumentType") or ""
    out.update({
        "Status": status,
        "Yahoo Name": meta.get("longName") or meta.get("shortName") or "",
        "Instrument Type": INSTRUMENT_TYPES.get(itype, itype.title()),
        "Exchange": meta.get("fullExchangeName") or meta.get("exchangeName") or "",
        "Currency": ccy,
        "Price Unit": unit,
        "Yahoo Currency Code": yccy,
        "First Date": df.index[0] if len(df) else None,
        "Last Date": last,
        "Daily Rows": len(df),
        "Notes": "; ".join(res["notes"]),
    })
    return out


NEW_COLUMNS = ["Yahoo Symbol", "Match Method", "Status", "Yahoo Name", "Instrument Type", "Exchange",
               "Currency", "Price Unit", "Yahoo Currency Code", "First Date", "Last Date", "Daily Rows", "Notes"]
WIDTHS = {"Yahoo Symbol": 14, "Match Method": 24, "Status": 26, "Yahoo Name": 30, "Instrument Type": 14,
          "Exchange": 14, "Currency": 10, "Price Unit": 14, "Yahoo Currency Code": 12, "First Date": 12,
          "Last Date": 12, "Daily Rows": 10, "Notes": 60}


def write_workbook(src, dst, rows):
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    table = next(iter(ws.tables.values()), None)
    first_new = ws.max_column + 1
    header_cell, body_cell = ws.cell(1, ws.max_column), ws.cell(2, ws.max_column)
    for j, col in enumerate(NEW_COLUMNS):
        c = ws.cell(1, first_new + j, col)
        c.font, c.fill, c.alignment, c.border = (copy(header_cell.font), copy(header_cell.fill),
                                                 copy(header_cell.alignment), copy(header_cell.border))
        ws.column_dimensions[get_column_letter(first_new + j)].width = WIDTHS[col]
    for i, row in enumerate(rows, start=2):
        for j, col in enumerate(NEW_COLUMNS):
            v = row.get(col)
            c = ws.cell(i, first_new + j, v if v not in (None, "") else None)
            c.font = copy(body_cell.font)
            if col in ("First Date", "Last Date") and v:
                c.number_format = "yyyy-mm-dd"
    if table is not None:
        table.ref = f"A1:{get_column_letter(first_new + len(NEW_COLUMNS) - 1)}{ws.max_row}"
        # Rebuild the column list so Excel doesn't flag the table as corrupt.
        table.tableColumns = [TableColumn(id=k + 1, name=str(ws.cell(1, k + 1).value))
                              for k in range(first_new + len(NEW_COLUMNS) - 1)]
        table.autoFilter.ref = table.ref
    ws.freeze_panes = "E2"
    wb.save(dst)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--out", required=True)
    ap.add_argument("--min-interval", type=float, default=2.0, help="Seconds between requests (default 2)")
    ap.add_argument("--limit", type=int, help="Only process the first N rows (for testing)")
    args = ap.parse_args()

    os.makedirs(os.path.join(args.out, "_state"), exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s",
                        handlers=[logging.StreamHandler(), logging.FileHandler(os.path.join(args.out, "_state", "run.log"))])
    os.makedirs(os.path.join(args.out, "prices"), exist_ok=True)
    state_path = os.path.join(args.out, "_state", "rows.json")
    state = json.load(open(state_path)) if os.path.exists(state_path) else {}

    src = pd.read_excel(args.xlsx)
    if args.limit:
        src = src.head(args.limit)
    yh = Yahoo(args.min_interval)
    today = dt.date.today()
    t0 = time.monotonic()

    for n, (_, r) in enumerate(src.iterrows(), 1):
        isin, ticker, name = r["ISIN"], r["Ticker"], r["Instrument Name"]
        if isin in state:
            continue
        res = resolve_row(yh, isin, ticker, name)
        info = describe(res, today)
        if res["symbol"] and res["df"] is not None and len(res["df"]):
            folder = os.path.join(args.out, "prices", res["symbol"])
            os.makedirs(folder, exist_ok=True)
            res["df"].to_csv(os.path.join(folder, "1d.csv"))
        state[isin] = {k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in info.items()}
        with open(state_path + ".tmp", "w") as f:
            json.dump(state, f, indent=1)
        os.replace(state_path + ".tmp", state_path)
        logger.info("[%d/%d] %-14s %-6s -> %-14s %-28s %s  (%d requests, %.0fs)", n, len(src), isin, ticker,
                    info["Yahoo Symbol"], info["Match Method"], info["Status"], yh.requests, time.monotonic() - t0)

    # Flag rows that resolve to the same Yahoo symbol (old/new ISINs, etc).
    rows = []
    for _, r in src.iterrows():
        row = dict(state[r["ISIN"]])
        for k in ("First Date", "Last Date"):
            if row.get(k):
                row[k] = dt.date.fromisoformat(row[k])
        rows.append(row)
    by_symbol = {}
    for i, row in enumerate(rows):
        if row.get("Yahoo Symbol"):
            by_symbol.setdefault(row["Yahoo Symbol"], []).append(i)
    for sym, idxs in by_symbol.items():
        if len(idxs) > 1:
            for i in idxs:
                others = [str(src.iloc[j]["No."]) for j in idxs if j != i]
                msg = f"Same Yahoo symbol as row No. {', '.join(others)} (likely old/new ISIN for one company)"
                rows[i]["Notes"] = "; ".join(x for x in (rows[i].get("Notes"), msg) if x)

    base = os.path.splitext(os.path.basename(args.xlsx))[0]
    base = re.sub(r"^[0-9a-f]{8}-", "", base)  # drop upload prefix
    dst = os.path.join(args.out, f"{base}_updated.xlsx")
    write_workbook(args.xlsx, dst, rows)
    print(f"\nWrote {dst}\n{yh.requests} requests in {time.monotonic() - t0:.0f}s")
    print(pd.Series([r["Status"].split(" since")[0] for r in rows]).value_counts().to_string())
    print(pd.Series([r["Match Method"] for r in rows]).value_counts().to_string())


if __name__ == "__main__":
    sys.exit(main())
